from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import List

from app.core.dependencies import get_current_user
from app.database.connection import get_db
from app.models.enums import UserRoleEnum
from app.models.task_master import TaskMaster
from app.models.user import User
from app.schemas.task_master import TaskResponse

router = APIRouter(prefix="/tasks", tags=["Query - Tasks"])


@router.get("/", response_model=List[TaskResponse])
def get_all_tasks(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(TaskMaster)
    if current_user.role == UserRoleEnum.SUPERVISOR:
        if current_user.department_id is None:
            raise HTTPException(status_code=403, detail="Supervisor is not linked to a department.")
        query = query.filter(
            (TaskMaster.department_id == current_user.department_id)
            | TaskMaster.department_id.is_(None)
        )
    return query.all()


@router.get("/excel-template")
def get_excel_template():
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from fastapi.responses import Response

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TaskMasterTemplate"

    headers = [
        "department_id",
        "department_name",
        "task_name",
        "description",
        "priority_category",
        "estimated_time",
        "task_type",
        "required_role",
        "required_skill",
        "suggested_start_time",
        "due_time",
        "requires_proof",
        "task_state"
    ]
    ws.append(headers)

    sample_rows = [
        ["", "Nursing", "Vital Signs Monitoring", "Record scheduled vital signs and escalate out-of-range readings to the responsible clinician.", "HIGH", 30, "RECURRING", "NURSE", "VITAL_SIGNS", "07:15", "08:00", "TRUE", "ACTIVE"],
        ["", "Pharmacy", "Controlled Medication Count", "Complete the required count with an authorized witness and document both staff identities.", "CRITICAL", 20, "RECURRING", "PHARMACIST", "CONTROLLED_MEDICATIONS", "09:00", "10:00", "TRUE", "ACTIVE"]
    ]
    for r in sample_rows:
        ws.append(r)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="00478D", end_color="00478D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    col_widths = [18, 22, 30, 60, 18, 16, 18, 18, 28, 22, 16, 18, 14]
    for i, w in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = w

    for row in ws.iter_rows(min_row=1, max_row=len(sample_rows)+1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers_response = {
        "Content-Disposition": "attachment; filename=task_master_template.xlsx"
    }
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_response
    )


@router.get("/{task_id}", response_model=TaskResponse)
def get_task(
    task_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(TaskMaster).filter(TaskMaster.task_id == task_id)
    if current_user.role == UserRoleEnum.SUPERVISOR:
        if current_user.department_id is None:
            raise HTTPException(status_code=403, detail="Supervisor is not linked to a department.")
        query = query.filter(
            (TaskMaster.department_id == current_user.department_id)
            | TaskMaster.department_id.is_(None)
        )
    task = query.first()

    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    return task
