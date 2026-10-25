from io import BytesIO

import openpyxl
from sqlalchemy import create_engine, inspect, text

from app.database.migrations import _migrate_task_master
from app.models.enums import TaskTypeEnum
from app.routes.query.tasks import get_excel_template
from app.schemas.task_master import TaskCreate


def test_task_template_normalizes_required_role():
    task = TaskCreate(
        task_name="Role-constrained task",
        task_type=TaskTypeEnum.RECURRING,
        required_role="  Nurse  ",
    )

    assert task.required_role == "NURSE"


def test_task_master_migration_adds_role_without_removing_existing_columns():
    engine = create_engine("sqlite://")
    try:
        with engine.begin() as connection:
            connection.execute(text(
                "CREATE TABLE task_master (task_id INTEGER PRIMARY KEY, task_name VARCHAR(100))"
            ))

        _migrate_task_master(engine, inspect(engine), engine.dialect.name)

        columns = {column["name"] for column in inspect(engine).get_columns("task_master")}
        assert {
            "task_id",
            "task_name",
            "required_role",
            "requires_proof",
            "suggested_start_time",
            "due_time",
        } <= columns
    finally:
        engine.dispose()


def test_task_template_excel_includes_role_column():
    response = get_excel_template()
    workbook = openpyxl.load_workbook(BytesIO(response.body), read_only=True)
    try:
        worksheet = workbook.active
        headers = next(worksheet.iter_rows(values_only=True))
        first_task = next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))
        assert {
            "required_role",
            "required_skill",
            "suggested_start_time",
            "due_time",
            "requires_proof",
        } <= set(headers)
        assert len(first_task) == len(headers)
    finally:
        workbook.close()
